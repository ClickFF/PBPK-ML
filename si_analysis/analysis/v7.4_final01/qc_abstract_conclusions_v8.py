#!/usr/bin/env python3
"""QC — the rewritten v8.0 ABSTRACT and CONCLUSIONS.

Asserts that nothing outside those four paragraphs moved, that the author's 61-entry reference list and every
EndNote field are untouched, that every quantity resolves against v8.0's OWN Results, and that the
upstream-to-downstream narrative is actually present rather than a list of numbers.

    .venv_pbpk/bin/python repo/claude_new_Sep21/analysis/v7.4_final01/qc_abstract_conclusions_v8.py
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document

S21 = Path(__file__).resolve().parents[2]
D2 = S21 / "manuscript" / "v7.4" / "final_01" / "draft_02"
SRC = D2 / "revised ACS JCIM v8.0.docx"
NEW = D2 / "v8_framing" / "v8.0_abstract_conclusions.docx"
HI = D2 / "v8_framing" / "v8.0_abstract_conclusions_highlighted.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

FORBIDDEN = ["improvement", "superior", "advantage", "strategy preference", "outperform",
             "better than", "proves", "demonstrates that", "uniformly superior", "modest numerical"]
R = []


def chk(n, ok, d):
    R.append((n, bool(ok), d))


def main():
    a, b = Document(str(SRC)), Document(str(NEW))
    pa, pb = a.paragraphs, b.paragraphs
    ta, tb = [p.text for p in pa], [p.text for p in pb]
    ca = next(i for i, x in enumerate(ta) if x.strip().upper() == "CONCLUSIONS")
    cb = next(i for i, x in enumerate(tb) if x.strip().upper() == "CONCLUSIONS")

    chk("1. Paragraphs 0-10 byte-identical", ta[:11] == tb[:11],
        "first divergence: %s" % next((i for i in range(11) if ta[i] != tb[i]), "none"))
    chk("2. INTRODUCTION through end of DISCUSSION byte-identical (%d paragraphs)" % (ca - 14),
        ta[14:ca] == tb[14:cb], "first divergence: %s"
        % next((k + 14 for k in range(ca - 14) if ta[14 + k] != tb[14 + k]), "none"))
    chk("3. ACKNOWLEDGMENTS onward byte-identical", ta[ca + 2:] == tb[cb + 3:],
        "%d trailing paragraphs" % len(ta[ca + 2:]))

    ra = [x for x in ta if re.match(r"^\(\d+\)\s", x.strip())]
    rb = [x for x in tb if re.match(r"^\(\d+\)\s", x.strip())]
    nums = [int(re.match(r"^\((\d+)\)", x.strip()).group(1)) for x in rb]
    chk("4. Reference list unchanged: %d entries, 1..%d unreordered" % (len(rb), len(rb)),
        ra == rb and nums == list(range(1, len(nums) + 1)), "identical" if ra == rb else "CHANGED")

    fa = sum(1 for p in pa for r in p._p.findall(".//" + W + "fldChar")
             if r.get(W + "fldCharType") == "begin")
    fb = sum(1 for p in pb for r in p._p.findall(".//" + W + "fldChar")
             if r.get(W + "fldCharType") == "begin")
    chk("5. Every EndNote field in the document survives", fa == fb, "%d before, %d after" % (fa, fb))

    ab = " ".join(tb[i] for i in (11, 12, 13))
    con = tb[cb + 1] + " " + tb[cb + 2]
    chk("6. Abstract within the JCIM 250-word cap", len(ab.split()) <= 250, "%d words" % len(ab.split()))
    chk("7. Abstract still three paragraphs", all(tb[i].strip() for i in (11, 12, 13)),
        "%d / %d / %d words" % tuple(len(tb[i].split()) for i in (11, 12, 13)))
    chk("8. Conclusions split into two paragraphs",
        bool(tb[cb + 1].strip()) and bool(tb[cb + 2].strip()),
        "%d + %d words (was one of %d)" % (len(tb[cb + 1].split()), len(tb[cb + 2].split()),
                                           len(ta[ca + 1].split())))

    # every quantity must be in v8.0's OWN Results, not merely in draft_02's
    res8 = "\n".join(ta[87:ca])
    q = set(re.findall(r"\d+\.\d+|\b\d{2,}\b", ab + con))
    missing = sorted(n for n in q if n not in res8)
    chk("9. Every quantity resolves against v8.0's own RESULTS", not missing,
        "unsourced: %s" % missing if missing else "all %d traced" % len(q))
    chk("10. The 41.5-53.7% range is deliberately avoided (absent from v8.0's Results)",
        "41.5" not in ab + con and "53.7" not in ab + con, "not used")

    # the narrative, not a pile of numbers
    chk("11. Abstract opens on the bottleneck, not the method",
        re.search(r"clearance \(CL.{0,12}\) is the hardest", tb[11]), "first sentence names CL_sys")
    chk("12. Abstract states the causal routing explicitly",
        "clearance governed downstream" in tb[12] and "tracked" in tb[12],
        "'clearance governed downstream' + per-compound routing")
    chk("13. Abstract names the design that established it",
        "substituting inputs one at a time" in tb[12], "scenario substitution stated")
    chk("14. Abstract states the contribution, not just results",
        "rather than another PK predictor" in tb[13], "contribution sentence present")
    chk("15. Conclusions carries the downstream finding (v8.0's had none)",
        "accuracy of the predicted clearance input" in con and "halved" in con,
        "clearance finding restored")
    chk("16. Conclusions keeps a significance claim and a scope limit",
        "regulatory and industrial adoption" in con and "curated Simcyp compound files" in con,
        "both present")

    low = (ab + " " + con).lower()
    hits = []
    for w in FORBIDDEN:
        for m in re.finditer(re.escape(w), low):
            lo = max(0, m.start() - 90)
            if re.search(r"\b(no|not|nor|rather than|does not|did not|without|neither)\b",
                         (ab + " " + con)[lo:m.end()], re.I):
                continue
            hits.append("%r -> %s" % (w, (ab + " " + con)[lo:m.end() + 40]))
    chk("17. No unqualified comparative claim", not hits,
        "%d hit(s): %s" % (len(hits), hits[:2]) if hits else "0 across %d terms" % len(FORBIDDEN))
    chk("18. Causal language stays hedged",
        "not detectably different" in ab and "points to" in con and "indicating" in ab,
        "association/indication wording retained")

    h = Document(str(HI))
    hr = [x.text for x in h.paragraphs if re.match(r"^\(\d+\)\s", x.text.strip())]
    chk("19. Highlighted twin keeps the same reference list", hr == ra, "identical")
    ins = sum(1 for i in (11, 12, 13) for r in h.paragraphs[i].runs if r.font.highlight_color is not None)
    dele = sum(1 for i in (11, 12, 13) for r in h.paragraphs[i].runs if r.font.strike)
    chk("20. Highlighted twin marks insertions and deletions", ins and dele,
        "%d added / %d removed runs in the Abstract" % (ins, dele))

    w_ = max(len(n) for n, _, _ in R)
    for n, ok, d in R:
        print("%-*s  %s  %s" % (w_, n, "PASS" if ok else "FAIL", d))
    bad = [n for n, ok, _ in R if not ok]
    print("\n%d/%d checks pass" % (len(R) - len(bad), len(R)))
    if bad:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
