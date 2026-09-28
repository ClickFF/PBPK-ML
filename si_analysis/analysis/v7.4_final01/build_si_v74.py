#!/usr/bin/env python3
"""Build the v7.4 Supporting Information document from the v4 SI source + the Round 4C frozen captions.

Stages
  1  read manuscript/v4/SI/src/body_si.md (structure, section titles, table legends, prose: unchanged)
  2  remap every figure to the Round 4C frozen numbering and replace its caption with the frozen text verbatim
  3  insert the six figures Round 4B demoted from the main text, and Section S8 for the two that fall after S23
  4  remap every cross-reference, rewrite the Contents line
  5  notation pass: placeholder forms (f_u, CL_sys, log10, R2 ...) -> real sub/superscripts, CLsys terminology
  6  inject the nine table fragments, apply emphasis.emphasize(), write the markdown
  7  render .docx (on the v7.2 SI template) and .pdf

Writes only inside manuscript/v7.4/SI/.

    .venv_pbpk/bin/python repo/claude_new_Sep21/analysis/v7.4_final01/build_si_v74.py
"""
from __future__ import annotations

import importlib.util
import re
import subprocess
import sys
from pathlib import Path

import pandas as pd

S21 = Path(__file__).resolve().parents[2]
V4SI = S21 / "manuscript" / "v4" / "SI"
SI = S21 / "manuscript" / "v7.4" / "SI"
SRC = SI / "src" / "body_si.md"
OUT_MD = SI / "Supplementary_Information_v74.md"
FROZEN = S21 / "manuscript" / "v7.4" / "final_01" / "sources" / "si" / "captions_final.md"
MAP = S21 / "manuscript" / "v7.4" / "final_01" / "sources" / "si" / "si_figure_numbering_map.csv"
SCRIPTS = S21 / "manuscript" / "v4" / "scripts"


def load(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


EMPH = load(SCRIPTS / "emphasis.py", "emphasis")

# ---------------------------------------------------------------- notation
# Applied after all remapping. Longest first; the guard stops a token that is already inside a tag.
NOTATION = [
    (r"CL_sys", "CL<sub>sys</sub>"), (r"CLsys", "CL<sub>sys</sub>"), (r"CL_int", "CL<sub>int</sub>"),
    (r"VD_ss", "VD<sub>ss</sub>"), (r"C_max", "C<sub>max</sub>"), (r"f_u", "f<sub>u</sub>"),
    # CLRbase is a literal Simcyp parameter name and is left as written
    (r"CLint", "CL<sub>int</sub>"),
    (r"CLR\b", "CL<sub>R</sub>"), (r"VDss", "VD<sub>ss</sub>"),
    (r"Cmax", "C<sub>max</sub>"), (r"\bFu\b", "f<sub>u</sub>"),
    (r"AUC0[–-]t", "AUC<sub>0–t</sub>"), (r"AUC₀₋ₜ", "AUC<sub>0–t</sub>"),
    (r"log₁₀", "log<sub>10</sub>"), (r"\blog10\b", "log<sub>10</sub>"),
    (r"\bR2\b", "R²"), (r"\bpKa\b", "pK<sub>a</sub>"), (r"\bKp\b", "K<sub>p</sub>"),
]
# bare CL meaning systemic clearance; "renal CL", "total CL" spelled out and CL<sub> are left alone
BARE_CL = re.compile(r"(?<![A-Za-z_<>/])CL(?![A-Za-z_<]|</)")


def notation(t: str) -> str:
    for pat, rep in NOTATION:
        t = re.sub(r"(?<![<>/\w])" + pat + r"(?![^<>]*</sub>)", rep, t)

    def cl(m):
        pre = t[max(0, m.start() - 14):m.start()]
        if re.search(r"(renal|Renal|separate|total)\s*$", pre):
            return "CL"
        return "CL<sub>sys</sub>"

    return BARE_CL.sub(cl, t)


# ---------------------------------------------------------------- frozen captions
def frozen_captions() -> dict[int, str]:
    out, cur, buf = {}, None, []
    for line in FROZEN.read_text().splitlines():
        m = re.match(r"^\*\*Figure S(\d+)\*\*", line)
        if m:
            if cur:
                out[cur] = " ".join(buf).strip()
            cur, buf = int(m.group(1)), []
            continue
        if cur and line.startswith("> "):
            buf.append(line[2:].strip())
    if cur:
        out[cur] = " ".join(buf).strip()
    return out


def numbering():
    """old v4/Round-4B stem -> (new number, new stem)."""
    m = pd.read_csv(MAP)
    return {r.old_name: (int(str(r.final_number).split()[-1].lstrip("S")), r.final_filename_stem)
            for r in m.itertuples()}


def main():
    caps = frozen_captions()
    num = numbering()
    assert len(caps) == 25 and len(num) == 25, (len(caps), len(num))
    t = (V4SI / "src" / "body_si.md").read_text()

    # old v4 figure number -> Round 4C number, taken from the map rather than discovered while
    # rewriting, so the cross-reference remap can run BEFORE the frozen captions are substituted.
    # Running it afterwards would remap the already-final numbers inside those captions.
    old_to_new = {}
    for base, (new_n, _) in num.items():
        m = re.match(r"FigureS(\d+)_", base)
        if m:
            old_to_new[int(m.group(1))] = new_n

    # ---- 4a. cross-references, on the v4 text, before any caption is replaced
    def xref(m):
        out = m.group(0)
        for o in sorted({int(x) for x in re.findall(r"\d+", out)}, reverse=True):
            if o in old_to_new:
                out = re.sub(r"\bS%d\b" % o, "S\x00%d\x00" % old_to_new[o], out)
        return out

    t = re.sub(r"Figures? S\d+(?:\s*(?:,|and|to|–|-|through)\s*S\d+)*", xref, t)
    t = t.replace("\x00", "")
    # the v4 SI cited main-text Figure 3 for the paired ML forest; that display is now Figure S5
    t = t.replace("compound-paired inference is reported in Figure 3 and Table S2",
                  "compound-paired inference is reported in Figure S5 and Table S2")

    # ---- 2. figure embeds: remap stem + number, substitute the frozen caption
    n_fig = 0

    def fig(m):
        nonlocal n_fig
        cap, path = m.group(1), m.group(2)
        stem = Path(path).stem
        base, part = (stem[:-6], stem[-6:]) if stem.endswith(("_part1", "_part2")) else (stem, "")
        new_n, new_stem = num[base]
        n_fig += 1
        if part == "_part2":                      # keep v4's "(continued)" wording, renumber only
            cap = re.sub(r"^FIGURE S\d+ \(continued\)", "FIGURE S%d (continued)" % new_n, cap)
        else:
            cap = caps[new_n]
        return "![%s](figures/%s%s.png)" % (cap, new_stem, part)

    t = re.sub(r"!\[(.*?)\]\((figures/[^)]+)\)", fig, t)
    assert n_fig == 24, n_fig                      # 19 figures, 5 of them with a second page

    # ---- 3. the six figures Round 4B moved into the SI
    def embed(n):
        return "![%s](figures/%s.png)" % (caps[n], num[[k for k, v in num.items() if v[0] == n][0]][1])

    after = {1: 2, 4: 5, 13: 14, 15: 16}           # insert new figure V right after existing figure K
    for k, v in after.items():
        anchor = re.search(r"!\[FIGURE S%d\..*?\)\n" % k, t, re.S)
        assert anchor, k
        t = t[:anchor.end()] + "\n" + embed(v) + "\n" + t[anchor.end():]

    s8 = ("## S8. Alternative-Scenario Presentations and Extended All-Predicted Material\n\n"
          "This section holds two displays that present analyses reported elsewhere in a different form: the "
          "input-substitution analysis of PBPK Group A as independent branches rather than as a cumulative "
          "ladder, and the goodness-of-fit and per-compound ranking panels of the all-predicted scenario that "
          "are not carried in main-text Figure 6.\n\n"
          + embed(24) + "\n\n" + embed(25) + "\n")
    t = t.rstrip() + "\n\n" + s8

    # ---- 4b. Contents
    contents = (
        "**Contents.** Section S1, ML benchmarking (Tables S1–S3, Figures S1–S5). "
        "Section S2, encoder exposure (Table S4). "
        "Section S3, PBPK scenario configuration and provenance (Tables S5, S6). "
        "Section S4, PBPK compound-level results (Tables S7–S9, Figures S6–S18; Figures S6–S10 span two pages "
        "each). Section S5, applicability domain (Table S10, Figure S19). "
        "Section S6, information-matched clearance-strategy comparison and sensitivity to retained renal "
        "clearance (Table S11, Figure S20). "
        "Section S7, exploratory mechanistic stratification by ECCS class and empirical descriptors "
        "(Tables S12, S13, Figures S21–S23). "
        "Section S8, alternative-scenario presentations and extended all-predicted material "
        "(Figures S24, S25). "
        "Abbreviations: Fu, plasma unbound fraction; CLsys, systemic clearance; CLint, intrinsic clearance; "
        "VDss, volume of distribution at steady state; FE, fold error; GMFE, geometric mean fold error; "
        "S+, ADMET Predictor (Simulations Plus) v11 prediction; DL–ML, ATFP-embedding plus machine-learning "
        "model (ML Group A); C–T, concentration–time; log-NRMSE, RMSE of log₁₀ concentration residuals divided "
        "by the observed log₁₀ range of the curve; CLR, renal clearance (Simcyp CLRbase); HL, Hodges–Lehmann "
        "estimate; ECCS, Extended Clearance Classification System.")
    t = re.sub(r"^\*\*Contents\.\*\*.*$", lambda _: contents, t, count=1, flags=re.M)

    # ---- 5. notation
    t = notation(t)

    SRC.parent.mkdir(parents=True, exist_ok=True)
    SRC.write_text(t)

    # ---- 6. inject tables, emphasize
    for key in re.findall(r"\{\{TABLE:([^}]+)\}\}", t):
        frag = notation((SI / "tables" / (key + ".md")).read_text().strip())
        t = t.replace("{{TABLE:%s}}" % key, frag)
    assert "{{" not in t and "S[x]" not in t
    OUT_MD.write_text(EMPH.emphasize(t))
    print("wrote %s  (%d words)" % (OUT_MD.relative_to(S21), len(OUT_MD.read_text().split())))

    # ---- 7. docx, with <sub>/<sup> rendered as real Word runs
    MD = load(SCRIPTS / "md2docx.py", "md2docx")
    base_add_runs = MD.add_runs

    def add_runs(par, text, size=None, base_bold=False, highlight=None):
        for seg in re.split(r"(<sub>.*?</sub>|<sup>.*?</sup>)", text, flags=re.S):
            if not seg:
                continue
            m = re.fullmatch(r"<(sub|sup)>(.*?)</\1>", seg, re.S)
            if m:
                r = par.add_run(m.group(2))
                r.font.subscript = m.group(1) == "sub"
                r.font.superscript = m.group(1) == "sup"
                r.bold = True if base_bold else None
                if size:
                    from docx.shared import Pt
                    r.font.size = Pt(size)
            else:
                base_add_runs(par, seg, size=size, base_bold=base_bold, highlight=highlight)

    MD.add_runs = add_runs
    MD.convert(OUT_MD, OUT_MD.with_suffix(".docx"), "si")
    print("wrote", OUT_MD.with_suffix(".docx").relative_to(S21))

    # ---- 7b. pdf
    RP = load(SCRIPTS / "render_pdf.py", "render_pdf")
    base_inline = RP.inline
    UNESCAPE = [("&lt;sub&gt;", "<sub>"), ("&lt;/sub&gt;", "</sub>"),
                ("&lt;sup&gt;", "<super>"), ("&lt;/sup&gt;", "</super>")]

    def inline(x):
        out = base_inline(x)
        for a, b in UNESCAPE:
            out = out.replace(a, b)
        return out

    RP.inline = inline

    # col_widths() sizes a column from its longest "word"; the <sub> tags are markup, not glyphs, so
    # measuring them inflates wide columns and squeezes narrow ones (a 3-digit n column wrapped mid-number).
    base_cw = RP.col_widths

    def col_widths(rows, total):
        clean = [[re.sub(r"</?su[bp]>", "", c) for c in r] for r in rows]
        w = list(base_cw(clean, total))
        # A very long unbreakable token in one column (a CSV filename) can starve a narrow one until a
        # 3-digit count wraps mid-number. Give every column the width its longest token needs, taking the
        # difference from the columns that have slack.
        need = []
        for j in range(len(clean[0])):
            longest = max((len(x) for r in clean for x in r[j].split()), default=1)
            need.append(min(longest * 5.0 + 8.0, total * 0.35))
        short = sum(max(0.0, n - x) for n, x in zip(need, w))
        if short > 0:
            surplus = [max(0.0, x - n) for n, x in zip(need, w)]
            pool = sum(surplus)
            if pool > 0:
                take = min(short, pool)
                w = [x - take * s / pool for x, s in zip(w, surplus)]
            w = [max(x, n) for x, n in zip(w, need)]
            k = total / sum(w)
            w = [x * k for x in w]
        return w

    RP.col_widths = col_widths
    sys.argv = ["render_pdf.py", str(OUT_MD), "Supporting Information — v7.4"]
    RP.main()


if __name__ == "__main__":
    main()
