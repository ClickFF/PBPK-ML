#!/usr/bin/env python3
"""v7.4 final_01 — highlighted reading copy of the Methods+Results revision.

Same surgical approach as build_final01_docx.py, but rendered for review rather than submission:

  edit      word-level diff, deleted words red and struck, inserted words highlighted yellow
  insert    whole paragraph highlighted yellow
  delete    paragraph retained, rendered red and struck, so nothing disappears silently
  relocate  as delete, with a yellow marker naming the SI destination
  keep      untouched

Because the marking comes from the disposition table rather than from a similarity matcher, every mark is
exact, and any mark outside METHODS/RESULTS would be a scope violation. Reuses word_diff() and docx_runs()
from manuscript/v4/scripts/build_highlighted_v4.py, so the highlighting convention matches the v4 artifacts.

    .venv_pbpk/bin/python repo/claude_new_Sep21/analysis/v7.4_final01/build_final01_highlight.py
"""
from __future__ import annotations

import copy
import csv
import importlib.util
import re
import shutil
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

S21 = Path(__file__).resolve().parents[2]
FINAL = S21 / "manuscript" / "v7.4" / "final_01"
DRAFT = FINAL / "draft"
BASE = FINAL / "sources" / "baseline" / "revised ACS JCIM v7.2.docx"
TABLES = FINAL / "sources" / "tables" / "TABLE_maintext.docx"
OUT = DRAFT / "v7.4_methods_results_highlighted.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MATH = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


def load(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


HL = load(S21 / "manuscript" / "v4" / "scripts" / "build_highlighted_v4.py", "build_highlighted_v4")
CLEAN = load(Path(__file__).with_name("build_final01_docx.py"), "build_final01_docx")
FIX = load(Path(__file__).with_name("table_fixes.py"), "table_fixes")


def clear(par):
    for ch in list(par._p):
        if ch.tag != W + "pPr":
            par._p.remove(ch)


def main():
    mb = CLEAN.blocks("body_methods.md")
    rb = CLEAN.blocks("body_results.md")
    bl = {"methods": mb, "results": rb}
    dispo = {int(r["para_idx"]): r for r in
             csv.DictReader((FINAL / "changelog" / "v72_paragraph_disposition.csv").open())}

    shutil.copy2(BASE, OUT)
    doc = Document(str(OUT))
    paras = doc.paragraphs
    body = doc.element.body
    original = [p.text for p in paras]

    marks = {"edit": 0, "insert": 0, "delete": 0, "relocate": 0}
    out_of_scope = []

    # 1. edited paragraphs: word-level diff
    for sec, mp in (("methods", CLEAN.METHODS_MAP), ("results", CLEAN.RESULTS_MAP)):
        for bi, pi, _ in mp:
            if pi is None or not dispo.get(pi, {}).get("disposition", "keep").startswith("edit"):
                continue
            new = bl[sec][bi]
            runs = HL.word_diff(original[pi], new)
            clear(paras[pi])
            for om in paras[pi]._p.findall(".//" + MATH + "oMath"):
                om.getparent().remove(om)
            HL.docx_runs(paras[pi], runs)
            marks["edit"] += 1
            if not (20 <= pi <= 129):
                out_of_scope.append(pi)

    # 2. deleted and relocated paragraphs: retained, struck through
    for pi in CLEAN.DELETE + CLEAN.RELOCATE:
        par = paras[pi]
        txt = original[pi]
        clear(par)
        for om in par._p.findall(".//" + MATH + "oMath"):
            om.getparent().remove(om)
        runs = [(txt, "del")]
        if pi in CLEAN.RELOCATE:
            runs.append(("  [relocated to Supporting Information Section S7]", "ins"))
            marks["relocate"] += 1
        else:
            marks["delete"] += 1
        HL.docx_runs(par, runs)
        if not (20 <= pi <= 129):
            out_of_scope.append(pi)

    # 3. inserted paragraphs, after their anchors
    for sec, mapping in CLEAN.INSERTS.items():
        for anchor, idxs in mapping.items():
            prev = paras[anchor]._p
            for bi in idxs:
                new = copy.deepcopy(prev)
                for ch in list(new):
                    if ch.tag != W + "pPr":
                        new.remove(ch)
                prev.addnext(new)
                HL.docx_runs(Paragraph(new, None), [(bl[sec][bi], "ins")])
                prev = new
                marks["insert"] += 1

    # 4. tables: 1-3 from the frozen file, Table 4 object removed (its struck title and note remain)
    frozen = Document(str(TABLES))
    tbls = [el for el in body.iterchildren() if el.tag == W + "tbl"]
    for ti, el in enumerate(tbls):
        act = CLEAN.TABLE_ACTION.get(ti)
        if act is None:
            el.getparent().remove(el)
        else:
            new = copy.deepcopy(frozen.tables[act]._tbl)
            el.addnext(new)
            el.getparent().remove(el)

    FIX.fix_tables(doc)
    doc.save(str(OUT))
    for k, v in marks.items():
        print("%-9s %3d paragraphs marked" % (k, v))
    print("marks outside METHODS/RESULTS: %d %s" % (len(out_of_scope), out_of_scope or "(scope clean)"))
    print("wrote", OUT.relative_to(S21))


if __name__ == "__main__":
    main()
