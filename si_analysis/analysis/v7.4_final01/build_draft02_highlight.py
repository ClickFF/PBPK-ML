#!/usr/bin/env python3
"""Highlighted Methods+Results showing exactly what draft_02 changed relative to draft_01.

Starts from the draft_01 .docx and applies the same edits draft_02 applies, but renders them as a
word-level diff: deleted words red and struck through, inserted words highlighted yellow. Because the marks
come from the transform itself rather than from a similarity matcher, every renumbered Figure S number shows
up individually and nothing is missed.

    .venv_pbpk/bin/python repo/claude_new_Sep21/analysis/v7.4_final01/build_draft02_highlight.py
"""
from __future__ import annotations

import copy
import importlib.util
import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.text.paragraph import Paragraph

S21 = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

FINAL = S21 / "manuscript" / "v7.4" / "final_01"
OUT = FINAL / "draft_02" / "maintext" / "v7.4_methods_results_draft02_highlighted.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MATH = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


def load(p, n):
    s = importlib.util.spec_from_file_location(n, p)
    m = importlib.util.module_from_spec(s)
    sys.modules[n] = m
    s.loader.exec_module(m)
    return m


HL = load(S21 / "manuscript" / "v4" / "scripts" / "build_highlighted_v4.py", "build_highlighted_v4")
MTX = load(HERE / "build_draft02_maintext.py", "build_draft02_maintext")


def clear(par):
    for ch in list(par._p):
        if ch.tag != W + "pPr":
            par._p.remove(ch)


def main():
    shutil.copy2(FINAL / "draft" / "v7.4_methods_results.docx", OUT)
    doc = Document(str(OUT))
    marks = {"renumbered": 0, "caption": 0, "inserted": 0}
    anchor_p = None
    anchor = "Correlation does not establish causation."

    # Only METHODS and RESULTS are in scope. The Discussion keeps its v7.2 text verbatim, including its
    # stale v1_run4 / Figure S4 pointer: renumbering that reference would make a dead pointer look valid.
    span = {}
    for i, q in enumerate(doc.paragraphs):
        u = q.text.strip().upper()
        if u in ("METHODS", "RESULTS", "DISCUSSION"):
            span.setdefault(u, i)
    lo, hi = span["METHODS"], span["DISCUSSION"]

    for i, p in enumerate(doc.paragraphs):
        if not (lo <= i < hi):
            continue
        old = p.text
        if not old.strip():
            continue
        if old.strip().startswith("FIGURE 2."):
            new = re.sub(r"</?su[bp]>|\*\*", "", MTX.FIG2_CAPTION)
            clear(p)
            HL.docx_runs(p, HL.word_diff(old, new))
            marks["caption"] += 1
            continue
        new = MTX.transform(old)
        if new != old:
            assert not p._p.findall(".//" + MATH + "oMath"), "equation paragraph"
            clear(p)
            HL.docx_runs(p, HL.word_diff(old, new))
            marks["renumbered"] += 1
        if anchor in p.text:
            anchor_p = p

    assert anchor_p is not None
    new_p = copy.deepcopy(anchor_p._p)
    for ch in list(new_p):
        if ch.tag != W + "pPr":
            new_p.remove(ch)
    anchor_p._p.addnext(new_p)
    HL.docx_runs(Paragraph(new_p, None),
                 [(re.sub(r"</?su[bp]>|\*\*", "", MTX.FIG6_CAPTION), "ins")])
    marks["inserted"] += 1
    doc.save(str(OUT))

    d = Document(str(OUT))
    sec = {}
    for i, q in enumerate(d.paragraphs):
        u = q.text.strip().upper()
        if u in ("METHODS", "RESULTS", "DISCUSSION"):
            sec.setdefault(u, i)
    out_of_scope = sum(1 for i, q in enumerate(d.paragraphs)
                       if any(r.font.highlight_color == WD_COLOR_INDEX.YELLOW or r.font.strike for r in q.runs)
                       and not (sec["METHODS"] <= i < sec["DISCUSSION"]))
    for k, v in marks.items():
        print("%-11s %3d paragraphs" % (k, v))
    print("marks outside METHODS/RESULTS: %d" % out_of_scope)
    print("wrote", OUT.relative_to(S21))


if __name__ == "__main__":
    main()
