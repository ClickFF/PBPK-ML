#!/usr/bin/env python3
"""Round 5 draft_02: rebuild the Supporting Information after the author's figure changes.

  * S15 + S16 leave the SI (they are now main-text Figure 6)
  * S24 moves into Section S4; Section S8 is dropped
  * S2, S20 and S25 are deleted as duplicates of main-text displays
  * the remaining 20 figures are renumbered S1-S20 and every cross-reference is remapped

Assets and documents are written under manuscript/v7.4/final_01/draft_02/.

    .venv_pbpk/bin/python repo/claude_new_Sep21/analysis/v7.4_final01/build_draft02_si.py
"""
from __future__ import annotations

import importlib.util
import re
import shutil
import sys
from pathlib import Path

S21 = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))
import draft02_map as M                                            # noqa: E402

V74 = S21 / "manuscript" / "v7.4"
SRC_SI = V74 / "SI"
D2 = V74 / "final_01" / "draft_02"
SI = D2 / "SI"
SCRIPTS = S21 / "manuscript" / "v4" / "scripts"


def load(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


EMPH = load(SCRIPTS / "emphasis.py", "emphasis")

# references that cannot simply be renumbered, because their target left the SI
SEMANTIC = [
    # Table S11 / Section S6 lost their figure
    ("(Table S11, Figure S20)", "(Table S11)"),
    # the S16 caption's closing pointer disappears with the figure itself
]


def copy_assets():
    (D2 / "figures").mkdir(parents=True, exist_ok=True)
    (SI / "figures").mkdir(parents=True, exist_ok=True)
    n_main = n_si = 0
    for stem in M.MAIN_KEEP:
        for ext in ("png", "pdf", "svg"):
            src = V74 / "figures_round4b" / ("%s.%s" % (stem, ext))
            if src.exists():
                shutil.copy2(src, D2 / "figures" / src.name)
                n_main += 1
    for old, (new, _) in M.MAIN_RENAME.items():
        for ext in ("png", "pdf", "svg"):
            src = V74 / "figures_round4b" / ("%s.%s" % (old, ext))
            if src.exists():
                shutil.copy2(src, D2 / "figures" / ("%s.%s" % (new, ext)))
                n_main += 1
    for stem, _, new_stem in M.SI_ORDER:
        for p in sorted((SRC_SI / "figures").glob(stem + "*")):
            suffix = p.stem[len(stem):]              # "" or "_part1"/"_part2"
            shutil.copy2(p, SI / "figures" / ("%s%s%s" % (new_stem, suffix, p.suffix)))
            n_si += 1
    print("copied %d main-text figure files, %d SI figure files" % (n_main, n_si))


def rebuild_si_source() -> str:
    t = (SRC_SI / "src" / "body_si.md").read_text()

    # 1. lift the S24 embed out, drop Section S8 and the removed figures
    embeds = re.findall(r"!\[.*?\]\(figures/[^)]+\)", t, re.S)
    s24 = [e for e in embeds if "FigureS24_" in e]
    assert len(s24) == 1
    t = t.replace(s24[0] + "\n\n", "").replace(s24[0], "")
    t = re.sub(r"\n## S8\..*$", "\n", t, flags=re.S)
    for num, (stem, why) in M.REMOVED.items():
        gone = [e for e in re.findall(r"!\[.*?\]\(figures/[^)]+\)", t, re.S) if stem in e]
        for e in gone:
            t = t.replace(e + "\n\n", "").replace(e, "")
        print("  removed Figure S%-2d (%s)" % (num, why))

    # 2. put S24 back, at the end of the scenario-error material in Section S4
    anchor = re.search(r"!\[[^\]]*FIGURE S14\..*?\)\n", t, re.S)
    assert anchor, "error-distribution figure not found"
    t = t[:anchor.end()] + "\n" + s24[0] + "\n" + t[anchor.end():]

    # 3. renumber captions and cross-references in one simultaneous pass
    def renum(m):
        out = m.group(0)
        for o in sorted({int(x) for x in re.findall(r"\d+", out)}, reverse=True):
            if o in M.NEW_NUM:
                out = re.sub(r"\bS%d\b" % o, "S\x00%d\x00" % M.NEW_NUM[o], out)
        return out

    t = re.sub(r"FIGURE S\d+|Figures? S\d+(?:\s*(?:,|and|to|–|-|through)\s*S\d+)*", renum, t)
    t = t.replace("\x00", "")
    for a, b in SEMANTIC:
        t = t.replace(a, b)

    # 4. image paths to the new stems
    for stem, _, new_stem in M.SI_ORDER:
        t = t.replace("figures/%s" % stem, "figures/%s" % new_stem)

    # 5. Contents
    contents = (
        "**Contents.** Section S1, ML benchmarking (Tables S1–S3, Figures S1–S4). "
        "Section S2, encoder exposure (Table S4). "
        "Section S3, PBPK scenario configuration and provenance (Tables S5, S6). "
        "Section S4, PBPK compound-level results (Tables S7–S9, Figures S5–S16; Figures S5–S9 span two pages "
        "each). Section S5, applicability domain (Table S10, Figure S17). "
        "Section S6, information-matched clearance-strategy comparison and sensitivity to retained renal "
        "clearance (Table S11). "
        "Section S7, exploratory mechanistic stratification by ECCS class and empirical descriptors "
        "(Tables S12, S13, Figures S18–S20). "
        "The association between upstream input error and downstream exposure error, previously Figures S15 "
        "and S16 of this document, is now Figure 6 of the main text. "
        "Abbreviations: Fu, plasma unbound fraction; CLsys, systemic clearance; CLint, intrinsic clearance; "
        "VDss, volume of distribution at steady state; FE, fold error; GMFE, geometric mean fold error; "
        "S+, ADMET Predictor (Simulations Plus) v11 prediction; DL–ML, ATFP-embedding plus machine-learning "
        "model (ML Group A); C–T, concentration–time; log-NRMSE, RMSE of log₁₀ concentration residuals divided "
        "by the observed log₁₀ range of the curve; CLR, renal clearance (Simcyp CLRbase); HL, Hodges–Lehmann "
        "estimate; ECCS, Extended Clearance Classification System.")
    BUILD = load(Path(__file__).with_name("build_si_v74.py"), "build_si_v74")
    contents = BUILD.notation(contents)      # the body is already converted; this string is not
    t = re.sub(r"^\*\*Contents\.\*\*.*$", lambda _: contents, t, count=1, flags=re.M)
    return re.sub(r"\n{3,}", "\n\n", t).rstrip() + "\n"


def main():
    copy_assets()
    t = rebuild_si_source()
    (SI / "src").mkdir(parents=True, exist_ok=True)
    (SI / "tables").mkdir(parents=True, exist_ok=True)
    for p in sorted((SRC_SI / "tables").glob("*.md")):
        shutil.copy2(p, SI / "tables" / p.name)
    (SI / "src" / "body_si.md").write_text(t)

    BUILD = load(Path(__file__).with_name("build_si_v74.py"), "build_si_v74")
    for key in re.findall(r"\{\{TABLE:([^}]+)\}\}", t):
        t = t.replace("{{TABLE:%s}}" % key,
                      BUILD.notation((SI / "tables" / (key + ".md")).read_text().strip()))
    assert "{{" not in t
    out_md = SI / "Supplementary_Information_v74b.md"
    out_md.write_text(EMPH.emphasize(t))
    print("wrote %s (%d words)" % (out_md.relative_to(S21), len(out_md.read_text().split())))

    MD = load(SCRIPTS / "md2docx.py", "md2docx")
    base_add_runs = MD.add_runs

    def add_runs(par, text, size=None, base_bold=False, highlight=None):
        for seg in re.split(r"(<sub>.*?</sub>|<sup>.*?</sup>)", text, flags=re.S):
            if not seg:
                continue
            m = re.fullmatch(r"<(sub|sup)>(.*?)</\1>", seg, re.S)
            if m:
                r = par.add_run(m.group(2))
                r.font.subscript = m.group(1) == "sub"
                r.font.superscript = m.group(1) == "sup"
                r.bold = True if base_bold else None
                if size:
                    from docx.shared import Pt
                    r.font.size = Pt(size)
            else:
                base_add_runs(par, seg, size=size, base_bold=base_bold, highlight=highlight)

    MD.add_runs = add_runs
    MD.convert(out_md, out_md.with_suffix(".docx"), "si")
    print("wrote", out_md.with_suffix(".docx").relative_to(S21))

    RP = load(SCRIPTS / "render_pdf.py", "render_pdf")
    base_inline, base_cw = RP.inline, RP.col_widths

    def inline(x):
        out = base_inline(x)
        for a, b in (("&lt;sub&gt;", "<sub>"), ("&lt;/sub&gt;", "</sub>"),
                     ("&lt;sup&gt;", "<super>"), ("&lt;/sup&gt;", "</super>")):
            out = out.replace(a, b)
        return out

    def col_widths(rows, total):
        clean = [[re.sub(r"</?su[bp]>", "", c) for c in r] for r in rows]
        w = list(base_cw(clean, total))
        need = []
        for j in range(len(clean[0])):
            longest = max((len(x) for r in clean for x in r[j].split()), default=1)
            need.append(min(longest * 5.0 + 8.0, total * 0.35))
        if sum(max(0.0, n - x) for n, x in zip(need, w)) > 0:
            surplus = [max(0.0, x - n) for n, x in zip(need, w)]
            pool = sum(surplus)
            if pool > 0:
                take = min(sum(max(0.0, n - x) for n, x in zip(need, w)), pool)
                w = [x - take * s / pool for x, s in zip(w, surplus)]
            w = [max(x, n) for x, n in zip(w, need)]
            k = total / sum(w)
            w = [x * k for x in w]
        return w

    RP.inline, RP.col_widths = inline, col_widths
    sys.argv = ["render_pdf.py", str(out_md), "Supporting Information — v7.4 (draft_02)"]
    RP.main()


if __name__ == "__main__":
    main()
