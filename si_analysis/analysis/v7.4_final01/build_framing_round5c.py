#!/usr/bin/env python3
"""Round 5C — Introduction, Abstract and Conclusions, built on the frozen Round 5C Discussion manuscript.

Input is discussion/v7.4_discussion_round5c.docx, whose Discussion the author froze on 2026-09-24. This
script changes only the Abstract, the Introduction and the Conclusions; Methods, Results, the frozen
Discussion and the whole reference list are left untouched.

Paragraphs 13-16 of the Introduction carry 35 live EndNote fields between them, so they are edited by
run-span substitution rather than rebuilt. The Abstract and the Conclusions carry no fields and are
rewritten in full.

    .venv_pbpk/bin/python repo/claude_new_Sep21/analysis/v7.4_final01/build_framing_round5c.py
"""
from __future__ import annotations

import copy
import importlib.util
import re
import shutil
import sys
from pathlib import Path

from docx import Document

S21 = Path(__file__).resolve().parents[2]
D2 = S21 / "manuscript" / "v7.4" / "final_01" / "draft_02"
SRC = D2 / "discussion" / "v7.4_discussion_round5c.docx"
OUT = D2 / "framing_draft"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
BD = None

# ----------------------------------------------------------------- INTRODUCTION (surgical; fields intact)
INTRO = [
    # VD_ss naming, consistent with the rest of the manuscript (R3 Minor 13)
    (13, "the volume of distribution (Vd) and total body clearance",
         "the volume of distribution at steady state (VD<sub>ss</sub>) and total body clearance"),
    (13, "As another fundamental PK property, Vd can often be predicted",
         "As another fundamental PK property, VD<sub>ss</sub> can often be predicted"),
    # the sentence fragment left after the citation
    (15, " , which is referred as quantitative structure activity relationship models (PK–QSAR).",
         " These are commonly referred to as quantitative structure–activity relationship models for PK "
         "endpoints (PK–QSAR)."),
    (15, "Many reported models rely on predefined descriptors  including",
         "Many reported models rely on predefined descriptors including"),
    # name the gap this study fills: the hinge to the upstream-accuracy result
    (16, "Nevertheless, most investigations only evaluate aggregate performance metrics and typically "
         "employ a single clearance parameterization strategy.",
         "Nevertheless, most investigations report aggregate performance metrics under a single clearance "
         "parameterization strategy, and rarely trace which of the predicted inputs drives the downstream "
         "exposure error."),
    # re-aim the second question onto what the Results actually answer
    (17, "Second, we examine how alternative clearance parameterizations and PBPK model complexity "
         "influence exposure predictions, both at the global level and for individual compounds, within a "
         "hybrid modeling framework.",
         "Second, we ask which predicted input governs the downstream exposure error, and whether the "
         "choice between two predicted-clearance representations changes that error once each receives "
         "the same information."),
    # the comma splice (R3 Minor 10); the two questions are already stated in 17, so the duplicate goes
    (19, " platform., we investigated (i) whether learned molecular embeddings can replace handcrafted "
         "descriptors without sacrificing predictive performance, and (ii) how bottom-up (IVIVE-scaled "
         "CLint) versus CLsys-based clearance parameterizations affect exposure predictions in hybrid "
         "PBPK modeling under standardized evaluation conditions.",
         " platform."),
]

# ----------------------------------------------------------------- ABSTRACT (rebuilt; no fields present)
ABSTRACT = (
    "Accurate prediction of human pharmacokinetics (PK) remains challenging in early drug discovery, "
    "particularly when experimentally determined ADME parameters are unavailable. We developed a "
    "graph-based deep learning encoder–machine learning (DL–ML) framework for key human PK parameters and "
    "evaluated its integration into physiologically based pharmacokinetic (PBPK) simulations in Simcyp. "
    "On a standardized benchmark dataset, graph attention–derived molecular embeddings combined with RDKit "
    "descriptors matched descriptor-based ML models for systemic clearance (CL<sub>sys</sub>), plasma "
    "unbound fraction (Fu) and volume of distribution at steady state (VD<sub>ss</sub>), with no detectable "
    "paired difference against reproduced descriptor models. For 41 compounds, we then "
    "compared bottom-up (IVIVE-scaled intrinsic clearance, CL<sub>int</sub>) and CL<sub>sys</sub>-based "
    "parameterizations under identical PBPK structures. Predicted CL<sub>sys</sub> was the least accurate "
    "structure-derived input (within twofold for 54% of compounds), and replacing observed clearance with a "
    "predicted value reduced AUC<sub>0–t</sub> twofold coverage from 82.9% to 41.5–53.7%; substituting "
    "physicochemical properties, Fu or VD<sub>ss</sub> changed exposure error little. Under matched "
    "renal-clearance information the two representations were not detectably different on any accuracy "
    "endpoint, which at this sample size does not establish equivalence. With Fu, VD<sub>ss</sub> and "
    "CL<sub>sys</sub> all predicted, AUC<sub>0–t</sub> was within twofold for 41% of compounds and "
    "C<sub>max</sub> for 59%. A systematic C<sub>max</sub> underprediction persisted under observed inputs, "
    "indicating a contribution from the minimal PBPK structure rather than input error. Downstream exposure "
    "accuracy was therefore governed primarily by the accuracy of the predicted clearance input; "
    "graph-based embeddings provide complementary structural information alongside conventional "
    "descriptors. The framework provides a structured basis for integrating structure-derived PK parameters "
    "into mechanistic exposure simulations."
)

# ----------------------------------------------------------------- CONCLUSIONS (rebuilt, split in two)
CONCLUSIONS = [
    "In summary, this study presents a structured hybrid ML–PBPK workflow for evaluating how "
    "structure-derived PK parameter predictions propagate within PBPK exposure simulations under "
    "standardized modeling conditions. Graph-based molecular embeddings achieved performance for key human "
    "PK parameter predictions broadly comparable to conventional descriptor-based ML models and provided "
    "complementary molecular information when integrated with RDKit descriptors. Within the hybrid PBPK "
    "framework, exposure accuracy was governed primarily by the accuracy of the predicted clearance input: "
    "replacing an observed clearance value with a predicted one roughly halved the proportion of compounds "
    "whose AUC<sub>0–t</sub> fell within twofold of the observed value, whereas the two predicted-clearance "
    "representations were not detectably different once matched for renal-clearance information. A residual "
    "C<sub>max</sub> underprediction that persisted under observed inputs is consistent with a contribution "
    "from the minimal PBPK structure, and escalation to full PBPK did not consistently improve accuracy "
    "when inputs were predicted.",

    "Hybrid ML–PBPK strategies have been explored in academic and open-source modeling environments. "
    "However, systematic comparison of alternative clearance parameterizations within a standardized and "
    "widely used commercial PBPK simulator has been limited. Given the regulatory and industrial adoption "
    "of platforms such as Simcyp, evaluating hybrid workflows under such environments is important for "
    "translational applicability. The present evaluation was conducted on compounds for which curated "
    "Simcyp compound files already exist, so it characterizes the propagation of predicted inputs through "
    "otherwise calibrated models; within that scope, the findings can inform practical guidance for "
    "applying ML/DL-informed ADME/PK parameter prediction within PBPK workflows and for identifying where "
    "such workflows are fit for purpose.",
]


def build(highlight: bool):
    out = OUT / ("v7.4_framing_round5c_highlighted.docx" if highlight else "v7.4_framing_round5c.docx")
    OUT.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC, out)
    doc = Document(str(out))
    paras = doc.paragraphs
    mode = "diff" if highlight else "clean"
    stats = {"intro": 0, "abstract": 0, "conclusions": 0, "moved": 0}

    missed = []
    for idx, old, new in INTRO:
        if not BD.rich_replace(paras[idx], old, new, mode):
            missed.append((idx, old[:55]))
        else:
            stats["intro"] += 1
    if missed:
        raise SystemExit("introduction anchors not found: %s" % missed)

    # the FIGURE 1 caption sits between the two statements of the research questions; it belongs after
    # the paragraph that introduces the figure
    cap = paras[18]
    if "FIGURE 1." not in cap.text:
        raise SystemExit("paragraph 18 is not the Figure 1 caption")
    paras[19]._p.addnext(cap._p)
    stats["moved"] = 1

    ab = paras[11]
    old_ab = ab.text
    for ch in list(ab._p):
        if ch.tag != W + "pPr":
            ab._p.remove(ch)
    if highlight:
        for t, kind in BD.HL.word_diff(old_ab, re.sub(r"</?sub>", "", ABSTRACT)):
            BD.HL.docx_runs(ab, [(t, kind)])
    else:
        BD.write(ab, ABSTRACT, {})
    stats["abstract"] = 1

    ci = next(i for i, p in enumerate(paras) if p.text.strip().upper() == "CONCLUSIONS")
    con = paras[ci + 1]
    old_con = con.text
    for ch in list(con._p):
        if ch.tag != W + "pPr":
            con._p.remove(ch)
    if highlight:
        for t, kind in BD.HL.word_diff(old_con, re.sub(r"</?sub>", "", " ".join(CONCLUSIONS))):
            BD.HL.docx_runs(con, [(t, kind)])
    else:
        BD.write(con, CONCLUSIONS[0], {})
        second = copy.deepcopy(con._p)
        for ch in list(second):
            if ch.tag != W + "pPr":
                second.remove(ch)
        con._p.addnext(second)
        from docx.text.paragraph import Paragraph
        BD.write(Paragraph(second, con._parent), CONCLUSIONS[1], {})
    stats["conclusions"] = 1 if highlight else 2

    doc.save(str(out))
    return out, stats



def write_md(clean):
    """Abstract, Introduction and Conclusions as markdown, read back from the built .docx."""
    doc = Document(str(clean))
    ps = doc.paragraphs
    ci = next(i for i, x in enumerate(ps) if x.text.strip().upper() == "CONCLUSIONS")
    spans = [("ABSTRACT", 10, 12), ("INTRODUCTION", 12, 20), ("CONCLUSIONS", ci, ci + 3)]

    out = ["<!-- Round 5C framing sections, read back from %s." % clean.name,
           "     Reference numbers are those of 'revised ACS JCIM v7.4.docx' and are not reordered.",
           "     Built by analysis/v7.4_final01/build_framing_round5c.py; QC: qc_framing_round5c_final.py.",
           "     The DISCUSSION is frozen and lives in discussion/v7.4_discussion_round5c.md. -->", ""]
    for name, lo, hi in spans:
        out += ["## " + name, ""]
        for i in range(lo, hi):
            p_ = ps[i]
            t = p_.text.strip()
            if not t or t.upper() == name:
                continue
            buf = []
            for r in p_.runs:
                if not r.text:
                    continue
                if r.font.subscript:
                    buf.append("<sub>%s</sub>" % r.text)
                elif r.font.superscript:
                    buf.append("<sup>%s</sup>" % r.text)
                else:
                    buf.append(r.text)
            line = "".join(buf)
            line = re.sub(r"</sub><sub>", "", line)
            line = re.sub(r"</sup><sup>", "", line)
            out += [line.strip(), ""]

    md = OUT / "v7.4_framing_round5c.md"
    md.write_text("\n".join(out))
    print("%-46s written" % md.name)


def main():
    global BD
    spec = importlib.util.spec_from_file_location(
        "bd", Path(__file__).with_name("build_discussion_round5c.py"))
    BD = importlib.util.module_from_spec(spec)
    sys.modules["bd"] = BD
    spec.loader.exec_module(BD)
    spec2 = importlib.util.spec_from_file_location(
        "hl", S21 / "manuscript" / "v4" / "scripts" / "build_highlighted_v4.py")
    BD.HL = importlib.util.module_from_spec(spec2)
    spec2.loader.exec_module(BD.HL)

    for hi in (False, True):
        out, stats = build(hi)
        print("%-46s %s" % (out.name, stats))

    write_md(OUT / "v7.4_framing_round5c.docx")
    d = Document(str(OUT / "v7.4_framing_round5c.docx"))
    print("\nabstract: %d words (JCIM cap 250)" % len(d.paragraphs[11].text.split()))
    ci = next(i for i, p in enumerate(d.paragraphs) if p.text.strip().upper() == "CONCLUSIONS")
    print("conclusions: %d + %d words, 2 paragraphs"
          % (len(d.paragraphs[ci + 1].text.split()), len(d.paragraphs[ci + 2].text.split())))


if __name__ == "__main__":
    main()
