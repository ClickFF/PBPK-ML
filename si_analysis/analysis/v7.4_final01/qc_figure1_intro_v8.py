#!/usr/bin/env python3
"""QC — the FIGURE 1 caption and the two-objective paragraph in v8.0_framing_full.docx."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document

S21 = Path(__file__).resolve().parents[2]
V = S21 / "manuscript" / "v7.4" / "final_01" / "draft_02" / "v8_framing"
SRC, NEW = V / "v8.0_abstract_conclusions.docx", V / "v8.0_framing_full.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = []


def chk(n, ok, d):
    R.append((n, bool(ok), d))


def main():
    a, b = Document(str(SRC)), Document(str(NEW))
    ta, tb = [p.text for p in a.paragraphs], [p.text for p in b.paragraphs]
    cap = next(i for i, x in enumerate(tb) if x.strip().startswith("FIGURE 1."))
    intro = next(i for i, x in enumerate(tb) if "as shown in Figure 1" in x)

    chk("1. Only the caption and the introducing paragraph changed",
        [i for i in range(len(ta)) if ta[i] != tb[i]] == [cap, intro],
        "changed paragraphs %s" % [i for i in range(len(ta)) if ta[i] != tb[i]])
    chk("2. Paragraph count unchanged (nothing inserted or dropped)",
        len(ta) == len(tb), "%d -> %d" % (len(ta), len(tb)))

    ra = [x for x in ta if re.match(r"^\(\d+\)\s", x.strip())]
    rb = [x for x in tb if re.match(r"^\(\d+\)\s", x.strip())]
    chk("3. Reference list untouched (%d entries)" % len(rb), ra == rb, "identical")
    fa = sum(1 for p in a.paragraphs for r in p._p.findall(".//" + W + "fldChar")
             if r.get(W + "fldCharType") == "begin")
    fb = sum(1 for p in b.paragraphs for r in p._p.findall(".//" + W + "fldChar")
             if r.get(W + "fldCharType") == "begin")
    chk("4. Every EndNote field survives", fa == fb, "%d before, %d after" % (fa, fb))

    capp = b.paragraphs[cap]
    chk("5. The inline FIGURE 1 image survived the caption rewrite",
        len(capp._p.findall(".//" + W + "drawing")) == 1, "one inline drawing present")
    chk("6. No literal <sub>/<sup> markup leaked",
        "<sub>" not in tb[cap] + tb[intro] and "<sup>" not in tb[cap] + tb[intro], "clean")
    chk("7. Subscripts and superscripts are real runs",
        {"sys", "ss", "max", "0–t"} <= {r.text for r in capp.runs if r.font.subscript}
        and any(r.font.superscript for r in capp.runs),
        "sub %s / sup %s" % (sorted({r.text for r in capp.runs if r.font.subscript}),
                             sorted({r.text for r in capp.runs if r.font.superscript})))

    chk("8. Caption walks all four panels of the figure",
        all(k in tb[cap] for k in ("(1A)", "(1B)", "(2A)", "(2B)")), "1A/1B/2A/2B all described")
    chk("9. Caption names both objectives as the figure labels them",
        "Objective 1, representation benchmarking" in tb[cap]
        and "Objective 2, hybrid PBPK modeling" in tb[cap], "both present")
    chk("10. Caption names both test sets", "Test #1" in tb[cap] and "Test #2" in tb[cap], "both named")
    chk("11. Caption reports no result", not re.search(r"\d+(\.\d+)?%|\bfold error of\b", tb[cap]),
        "descriptive only, no numbers")

    chk("12. The comma splice is repaired", "platform., we investigated" not in tb[intro], "fixed")
    chk("13. The two objectives are restored in the text",
        "two objectives" in tb[intro] and "representation benchmarking" in tb[intro]
        and "hybrid PBPK modeling" in tb[intro], "both stated, matching the figure")
    q = next(i for i, x in enumerate(tb) if "Two methodological questions" in x)
    chk("14. Objectives are worded as objectives, not as a second copy of the questions",
        "we ask" not in tb[intro] and "(i) whether" not in tb[intro]
        and "we ask" in tb[q], "questions stay in ¶%d, objectives in ¶%d" % (q, intro))
    chk("15. Figure numbering untouched: FIGURE 1 is still FIGURE 1",
        sorted({int(m) for m in re.findall(r"FIGURE (\d)\.", "\n".join(tb))})
        == sorted({int(m) for m in re.findall(r"FIGURE (\d)\.", "\n".join(ta))}),
        "main-text figure numbers unchanged")

    w = max(len(n) for n, _, _ in R)
    for n, ok, d in R:
        print("%-*s  %s  %s" % (w, n, "PASS" if ok else "FAIL", d))
    bad = [n for n, ok, _ in R if not ok]
    print("\n%d/%d checks pass" % (len(R) - len(bad), len(R)))
    if bad:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
