#!/usr/bin/env python3
"""v8.0 — ABSTRACT and CONCLUSIONS rewritten around the upstream-to-downstream narrative.

Input is the author's `revised ACS JCIM v8.0.docx` (61 references, EndNote numbering of record). Only the
three Abstract paragraphs and the Conclusions are touched; every other paragraph, including the whole
Discussion and the reference list, is left byte-identical. None of the rewritten paragraphs contains an
EndNote field, so nothing is at risk there and no citation is added.

The narrative, in order: clearance is the hardest upstream parameter -> scenario-controlled substitution
shows it is the input that governs downstream exposure error -> the two clearance parameterizations are
indistinguishable once information-matched, so accuracy rather than representation is limiting -> a residual
C_max offset sits outside the inputs, in the distribution model -> the contribution is the controlled
end-to-end localization, not another predictor.

Every quantity used is present in v8.0's own RESULTS. 41.5-53.7% is deliberately NOT used: that range is in
draft_02's Results but not in v8.0's, so the narrative carries 82.9% -> 41% instead, which is.

    .venv_pbpk/bin/python repo/claude_new_Sep21/analysis/v7.4_final01/build_abstract_conclusions_v8.py
"""
from __future__ import annotations

import copy
import importlib.util
import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

S21 = Path(__file__).resolve().parents[2]
D2 = S21 / "manuscript" / "v7.4" / "final_01" / "draft_02"
SRC = D2 / "revised ACS JCIM v8.0.docx"
OUT = D2 / "v8_framing"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
BD = None

# --- ABSTRACT: three paragraphs, matching v8.0's own structure -----------------------------------------
ABSTRACT = [
    # setup and the ML layer: the bottleneck is named in the first sentence
    "Human pharmacokinetic (PK) prediction remains difficult in early drug discovery when experimental ADME "
    "parameters are unavailable, and systemic clearance (CL<sub>sys</sub>) is the hardest to predict. We "
    "developed a graph-based deep learning encoder–machine learning (DL–ML) framework for human PK "
    "parameters, then asked how its predictions behave as inputs to "
    "physiologically based pharmacokinetic (PBPK) simulation in Simcyp. On benchmark datasets, graph "
    "attention–derived embeddings with RDKit descriptors performed comparably to descriptor-based "
    "models for CL<sub>sys</sub>, unbound fraction (Fu) and steady-state volume of distribution "
    "(VD<sub>ss</sub>), with no detectable paired difference against reimplemented controls.",

    # the core: which input limits the simulation, and how that was established
    "Because parameter-level accuracy does not show which input limits a simulation, we evaluated 41 "
    "compounds across prespecified PBPK scenarios substituting inputs one at a time in a common "
    "structure. CL<sub>sys</sub> was the least accurate structure-derived input, within twofold for 54% of "
    "compounds, and clearance governed downstream performance: substituting physicochemical properties, Fu "
    "or VD<sub>ss</sub> changed exposure accuracy little, whereas replacing an observed clearance with a "
    "predicted one roughly halved twofold coverage of the area under the concentration–time curve "
    "(AUC<sub>0–t</sub>), from 82.9% to 41%. Per-compound errors located this routing: CL<sub>sys</sub> "
    "error tracked AUC<sub>0–t</sub>, and VD<sub>ss</sub> error tracked peak concentration "
    "(C<sub>max</sub>, within twofold for 59%).",

    # what the inputs do not explain, and the contribution
    "Bottom-up scaled intrinsic clearance (CL<sub>int</sub>) and predicted CL<sub>sys</sub> were not "
    "detectably different once matched for renal-clearance information, indicating that clearance accuracy, "
    "not its parameterization, is binding, while a residual C<sub>max</sub> underprediction persisted under "
    "observed inputs and varied with distribution-model structure. This work contributes a controlled, "
    "end-to-end localization of where structure-derived predictions constrain mechanistic exposure "
    "simulation, rather than another PK predictor.",
]

# --- CONCLUSIONS: two paragraphs, findings then significance and scope ----------------------------------
CONCLUSIONS = [
    "This study evaluated how structure-derived PK parameter predictions propagate into mechanistic exposure "
    "simulation under standardized modeling conditions. Graph-based molecular embeddings performed "
    "comparably to conventional descriptor-based ML models and provided complementary molecular information "
    "when combined with RDKit descriptors, but neither the choice of molecular representation nor the choice "
    "of clearance parameterization was what limited downstream performance: the accuracy of the predicted "
    "clearance input was. Scenario-controlled substitution localized the effect, in that non-clearance "
    "inputs could be replaced with little loss of exposure accuracy whereas predicting clearance "
    "approximately halved twofold AUC<sub>0–t</sub> coverage, and the two clearance parameterizations were "
    "indistinguishable once matched for renal-clearance information. A residual C<sub>max</sub> "
    "underprediction that persisted under observed inputs and varied with distribution-model structure "
    "points to a second source of error that lies in the model structure rather than in the predicted "
    "inputs.",

    "Systematic comparison of alternative clearance parameterizations within a standardized and widely used "
    "commercial PBPK simulator has been limited, and given the regulatory and industrial adoption of "
    "platforms such as Simcyp this evaluation is directly relevant to translational practice. Read together, "
    "the results indicate where effort is best placed in hybrid ML–PBPK workflows: on the accuracy of the "
    "predicted clearance input, and on the distribution model as a separate determinant of C<sub>max</sub>. "
    "Because the PBPK evaluation used compounds for which curated Simcyp compound files already exist, it "
    "characterizes the propagation of predicted inputs through otherwise calibrated models, and it defines "
    "the scope within which such workflows are currently fit for purpose.",
]


def rebuild(par, text, mode, old):
    for ch in list(par._p):
        if ch.tag != W + "pPr":
            par._p.remove(ch)
    if mode == "diff":
        for t, kind in BD.HL.word_diff(old, re.sub(r"</?sub>", "", text)):
            BD.HL.docx_runs(par, [(t, kind)])
    else:
        BD.write(par, text, {})


def add_after(par, text, mode):
    el = copy.deepcopy(par._p)
    for ch in list(el):
        if ch.tag != W + "pPr":
            el.remove(ch)
    par._p.addnext(el)
    np_ = Paragraph(el, par._parent)
    if mode == "diff":
        from docx.enum.text import WD_COLOR_INDEX
        for t in re.split(r"(<sub>.*?</sub>)", text):
            if not t:
                continue
            m = re.fullmatch(r"<sub>(.*?)</sub>", t)
            r = np_.add_run(m.group(1) if m else t)
            if m:
                r.font.subscript = True
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    else:
        BD.write(np_, text, {})
    return np_


def build(highlight: bool):
    out = OUT / ("v8.0_abstract_conclusions_highlighted.docx" if highlight
                 else "v8.0_abstract_conclusions.docx")
    OUT.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC, out)
    doc = Document(str(out))
    ps = doc.paragraphs
    mode = "diff" if highlight else "clean"

    for k, i in enumerate((11, 12, 13)):
        if sum(1 for r in ps[i]._p.findall(".//" + W + "fldChar")):
            raise SystemExit("paragraph %d unexpectedly carries a field" % i)
        rebuild(ps[i], ABSTRACT[k], mode, ps[i].text)

    ci = next(i for i, p in enumerate(ps) if p.text.strip().upper() == "CONCLUSIONS")
    con = ps[ci + 1]
    if sum(1 for r in con._p.findall(".//" + W + "fldChar")):
        raise SystemExit("the Conclusions paragraph unexpectedly carries a field")
    old = con.text
    if highlight:
        rebuild(con, " ".join(CONCLUSIONS), mode, old)
    else:
        rebuild(con, CONCLUSIONS[0], mode, old)
        add_after(con, CONCLUSIONS[1], mode)

    doc.save(str(out))
    return out


def main():
    global BD
    spec = importlib.util.spec_from_file_location(
        "bd", Path(__file__).with_name("build_discussion_round5c.py"))
    BD = importlib.util.module_from_spec(spec)
    sys.modules["bd"] = BD
    spec.loader.exec_module(BD)
    s2 = importlib.util.spec_from_file_location(
        "hl", S21 / "manuscript" / "v4" / "scripts" / "build_highlighted_v4.py")
    BD.HL = importlib.util.module_from_spec(s2)
    s2.loader.exec_module(BD.HL)

    for hi in (False, True):
        print("%-48s written" % build(hi).name)

    d = Document(str(OUT / "v8.0_abstract_conclusions.docx"))
    ab = " ".join(d.paragraphs[i].text for i in (11, 12, 13))
    ci = next(i for i, p in enumerate(d.paragraphs) if p.text.strip().upper() == "CONCLUSIONS")
    print("\nabstract   : %d words in 3 paragraphs (JCIM cap 250)" % len(ab.split()))
    print("conclusions: %d + %d words in 2 paragraphs"
          % (len(d.paragraphs[ci + 1].text.split()), len(d.paragraphs[ci + 2].text.split())))


if __name__ == "__main__":
    main()
