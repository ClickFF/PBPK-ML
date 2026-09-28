#!/usr/bin/env python3
"""Round 5 draft_02: update the Methods+Results text for the new figure architecture.

  * the merged parameter-to-exposure display becomes main-text Figure 6
  * the former Figure 6 (all-predicted workflow) becomes Figure 7
  * every SI cross-reference is remapped to the new S1-S20 numbering
  * references to the five deleted/moved SI figures are re-pointed, not renumbered

The .docx is produced by editing final_01/draft/v7.4_methods_results.docx paragraph by paragraph, so the
Word equation objects, EndNote fields and superscript citations of the untouched paragraphs survive.

    .venv_pbpk/bin/python repo/claude_new_Sep21/analysis/v7.4_final01/build_draft02_maintext.py
"""
from __future__ import annotations

import copy
import importlib.util
import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

S21 = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))
import draft02_map as M                                            # noqa: E402

FINAL = S21 / "manuscript" / "v7.4" / "final_01"
D2 = FINAL / "draft_02"
MT = D2 / "maintext"
SCRIPTS = S21 / "manuscript" / "v4" / "scripts"
def load(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MATH = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"

FIG6_CAPTION = (
    "**FIGURE 6. Association between upstream input error and downstream exposure error in the all-predicted "
    "scenario (DL–ML f<sub>u</sub>, VD<sub>ss</sub> and CL<sub>sys</sub> with S+ physicochemical inputs; "
    "h2_run0).** **(a)** Absolute log₂ fold error of the DL–ML CL<sub>sys</sub> prediction against the C–T "
    "RMSE of log₁₀ concentrations (N = 41); **(b)** signed log₂ CL<sub>sys</sub> fold error against signed "
    "log₂ AUC₀₋ₜ fold error; **(c)** signed log₂ VD<sub>ss</sub> fold error against signed log₂ "
    "C<sub>max</sub> fold error (predicted/observed; negative = underprediction). Each point is a compound; "
    "Spearman ρ with a 95% percentile bootstrap interval (10,000 resamples) and the two-sided p are given per "
    "panel. **(d)** Spearman correlation between each DL–ML input's log₂ fold error and each exposure error; "
    "colour encodes |ρ| and the printed value carries its sign; n.s. marks p ≥ 0.05 (n = 41; 40 for "
    "log-NRMSE). As a sensitivity analysis on the profile statistic, the association of the absolute "
    "CL<sub>sys</sub> fold error with the range-normalized C–T log-NRMSE (n = 40) is ρ = 0.71 [0.46, 0.87], "
    "p < 0.001. Associations are correlational; correlation does not establish causation.")

FIG2_CAPTION = (
    "**FIGURE 2. Dataset distribution and DL–ML model performance for the three PK parameters (ML Group "
    "A).** Rows are f<sub>u</sub>, CL<sub>sys</sub> and VD<sub>ss</sub>. **(a, d, g)** Distribution of the "
    "observed log₁₀ values in the training set (grey) and the held-out benchmark set (Test #1, blue), with "
    "the number of compounds, the full observed span and the span of the central 90%; the two distributions "
    "overlap closely, so the held-out set covers the same dynamic range as the training set. **(b, e, h)** "
    "Observed versus predicted values on the training set and **(c, f, i)** on the held-out set. Solid line, "
    "identity; shaded band, twofold; dotted lines, threefold; the observed axis is shared across the three "
    "panels of a row, and each scatter panel reports N, R² of the log₁₀ values, RMSE and the percentage "
    "within twofold. Training sets comprise 4042 f<sub>u</sub> records and 1287 each for CL<sub>sys</sub> "
    "and VD<sub>ss</sub>, of which 1284 are plotted for each of the latter two; three records fall outside "
    "the plotted range. Held-out sets comprise 633 f<sub>u</sub> and 177 each for CL<sub>sys</sub> and "
    "VD<sub>ss</sub>. [AUTHOR CHECK — caption drafted in Round 5; not part of the Round 4C freeze.]")

# references whose target left the SI, handled before the numeric remap
SEMANTIC = [
    # Figure 2 regains its distribution panels, so the paragraph cites them again; the numbers are the
    # measured spans, because the figure now shows them (the v7.2 "two to three orders" described the bulk)
    ("As summarized in TABLE 1, the training and test datasets covered wide dynamic ranges for all three PK "
     "parameters, spanning several orders of magnitude for Fu and approximately two to three orders for "
     "CLsys and VDss.",
     "As displayed in Figure 2a, d and g, the training and held-out sets covered wide and closely "
     "overlapping dynamic ranges, spanning 3.3 log₁₀ units for Fu, 5.5 for CLsys and 5.7 for VDss, with the "
     "central 90% of compounds within 2.1 to 2.5 units."),
    ("reported in Tables S8 and S11 and Figure S20.", "reported in Tables S8 and S11."),
    ("The renal-clearance sensitivity analysis is Figure S20, and the secondary comparison of the workflows "
     "as practically configured is in Table S11.",
     "The renal-clearance sensitivity analysis and the secondary comparison of the workflows as practically "
     "configured are reported in Table S11."),
    ("has to neutralize (Figure S20; Table S11)", "has to neutralize (Table S11)"),
    ("not detectably associated with any exposure endpoint (Figures S15 and S16)",
     "not detectably associated with any exposure endpoint (Figure 6)"),
    ("their relationship to upstream clearance prediction error is examined in the Supporting Information "
     "(Figures S15 and S16).",
     "their relationship to upstream clearance prediction error is examined in Figure 6."),
    (", and the extended goodness-of-fit and ranking panels are Figure S25.", "."),
    ("; the extended goodness-of-fit and ranking panels are Figure S25.", "."),
]


def transform(t: str) -> str:
    """Renumber the old Figure 6 to 7, re-point moved references, then remap SI numbering."""
    t = re.sub(r"\bFIGURE 6\.", "FIGURE 7.", t)
    t = re.sub(r"\bFigure 6\b", "Figure 7", t)
    for a, b in SEMANTIC:
        t = t.replace(a, b)

    def renum(m):
        out = m.group(0)
        for o in sorted({int(x) for x in re.findall(r"S(\d+)", out)}, reverse=True):
            if o in M.NEW_NUM:
                out = re.sub(r"\bS%d\b" % o, "S\x00%d\x00" % M.NEW_NUM[o], out)
        return out

    t = re.sub(r"Figures? S\d+(?:\s*(?:,|and|to|–|-|through)\s*S\d+)*", renum, t)
    return t.replace("\x00", "")


def main():
    MT.mkdir(parents=True, exist_ok=True)

    # ---- markdown
    for name in ("body_methods.md", "body_results.md"):
        (MT / name).write_text(transform((FINAL / "src" / name).read_text()))
    res = (MT / "body_results.md").read_text()
    res = re.sub(r"\*\*FIGURE 2\..*?(?=\n\n)", lambda _: FIG2_CAPTION, res, count=1, flags=re.S)
    (MT / "body_results.md").write_text(res)
    anchor = "Correlation does not establish causation."
    assert res.count(anchor) == 1
    res = res.replace(anchor, anchor + "\n\n" + FIG6_CAPTION, 1)
    (MT / "body_results.md").write_text(res)
    print("markdown: Figure 6 caption inserted after the propagation paragraph")

    # assembled, caption-resolved markdown (the bodies keep <<<FROZEN_CAPTION:n>>> markers)
    BUILD = load(Path(__file__).with_name("build_final01.py"), "build_final01")
    caps = BUILD.frozen_captions()
    caps[7] = transform(caps[6])                    # the old Figure 6 caption, renumbered and re-pointed
    parts = []
    for name in ("body_methods.md", "body_results.md"):
        txt = (MT / name).read_text()
        txt = re.sub(r"<<<FROZEN_CAPTION:(\d+)>>>",
                     lambda m: caps[7 if int(m.group(1)) == 6 else int(m.group(1))], txt)
        assert "<<<" not in txt
        parts.append(txt.rstrip() + "\n")
    (MT / "v7.4_methods_results_draft02.md").write_text("\n".join(parts))
    print("wrote", (MT / "v7.4_methods_results_draft02.md").relative_to(S21))

    # ---- docx, edited in place from the final_01 build
    out = MT / "v7.4_methods_results_draft02.docx"
    shutil.copy2(FINAL / "draft" / "v7.4_methods_results.docx", out)
    doc = Document(str(out))
    MD_ = importlib.util.spec_from_file_location("md2docx", SCRIPTS / "md2docx.py")
    MD = importlib.util.module_from_spec(MD_)
    sys.modules["md2docx"] = MD
    MD_.loader.exec_module(MD)

    def add_rich(par, text, base_bold=False):
        for seg in re.split(r"(<sub>.*?</sub>|<sup>.*?</sup>)", text, flags=re.S):
            if not seg:
                continue
            m = re.fullmatch(r"<(sub|sup)>(.*?)</\1>", seg, re.S)
            if m:
                r = par.add_run(m.group(2))
                r.font.subscript = m.group(1) == "sub"
                r.font.superscript = m.group(1) == "sup"
                r.bold = True if base_bold else None
            else:
                MD.add_runs(par, seg, base_bold=base_bold)

    # Only METHODS and RESULTS are in scope. The Discussion keeps its v7.2 text verbatim, including its
    # stale v1_run4 / Figure S4 pointer: renumbering that reference would make a dead pointer look valid.
    span = {}
    for i, q in enumerate(doc.paragraphs):
        u = q.text.strip().upper()
        if u in ("METHODS", "RESULTS", "DISCUSSION"):
            span.setdefault(u, i)
    lo, hi = span["METHODS"], span["DISCUSSION"]

    n_edit, anchor_p = 0, None
    for i, p in enumerate(doc.paragraphs):
        if not (lo <= i < hi):
            continue
        old = p.text
        if not old.strip():
            continue
        new = transform(old)
        if new != old:
            assert not p._p.findall(".//" + MATH + "oMath"), "would destroy an equation: %r" % old[:60]
            for r in list(p._p.findall(W + "r")) + list(p._p.findall(W + "hyperlink")):
                p._p.remove(r)
            add_rich(p, new)
            n_edit += 1
        if p.text.strip().startswith("FIGURE 2."):
            for r in list(p._p.findall(W + "r")) + list(p._p.findall(W + "hyperlink")):
                p._p.remove(r)
            add_rich(p, FIG2_CAPTION)
            n_edit += 1
        if anchor in p.text:
            anchor_p = p
    assert anchor_p is not None, "propagation paragraph not found in the .docx"

    new_p = copy.deepcopy(anchor_p._p)
    for ch in list(new_p):
        if ch.tag != W + "pPr":
            new_p.remove(ch)
    anchor_p._p.addnext(new_p)
    add_rich(Paragraph(new_p, None), FIG6_CAPTION)
    doc.save(str(out))

    d = Document(str(out))
    figs = sorted({int(n) for n in re.findall(r"(?<!S)\bFigure (\d+)\b", "\n".join(q.text for q in d.paragraphs))})
    sis = sorted({int(n) for n in re.findall(r"Figures? S(\d+)", "\n".join(q.text for q in d.paragraphs))})
    print("docx: %d paragraphs edited, Figure 6 caption inserted" % n_edit)
    print("      main-text figures cited: %s" % figs)
    print("      SI figures cited: %s" % sis)
    print("wrote", out.relative_to(S21))


if __name__ == "__main__":
    main()
