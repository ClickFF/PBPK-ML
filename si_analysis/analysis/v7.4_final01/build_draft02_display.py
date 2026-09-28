#!/usr/bin/env python3
"""Round 5 draft_02: paste-ready display documents, Figures 1-7 and Tables 1-3.

Same template and house style as the final_01 display documents (TABLE_maintext.docx supplies the styles,
page setup and fonts), updated for the new figure architecture:
  Figure 6 is the merged parameter-to-exposure display; the former Figure 6 is now Figure 7.

    .venv_pbpk/bin/python repo/claude_new_Sep21/analysis/v7.4_final01/build_draft02_display.py
"""
from __future__ import annotations

import copy
import importlib.util
import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from PIL import Image

S21 = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import draft02_map as M                                            # noqa: E402

FINAL = S21 / "manuscript" / "v7.4" / "final_01"
D2 = FINAL / "draft_02"
OUT = D2 / "display"
FIGDIR = D2 / "figures"
TEMPLATE = FINAL / "sources" / "tables" / "TABLE_maintext.docx"
BUILT_MAIN = D2 / "maintext" / "v7.4_methods_results_draft02.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MAX_W_IN, MAX_H_IN = 6.5, 7.2


def load(p, n):
    s = importlib.util.spec_from_file_location(n, p)
    m = importlib.util.module_from_spec(s)
    sys.modules[n] = m
    s.loader.exec_module(m)
    return m


DISP = load(HERE / "build_display_docs.py", "build_display_docs")
FIX = load(HERE / "table_fixes.py", "table_fixes")
MTX = load(HERE / "build_draft02_maintext.py", "build_draft02_maintext")


def captions() -> dict[int, str]:
    """Figures 1-5 and 7 keep their approved captions; 6 is the merged one."""
    caps = {1: DISP.FIG1_TITLE}
    body = (D2 / "maintext" / "body_results.md").read_text()
    caps[2] = re.search(r"^(\*\*FIGURE 2\..*?)$", body, re.M).group(1).strip()
    frozen = DISP.BUILD.frozen_captions()          # Round 4C, Figures 3-6 (old numbering)
    for n in (3, 4, 5):
        caps[n] = frozen[n]
    caps[6] = MTX.FIG6_CAPTION
    caps[7] = MTX.transform(frozen[6])             # old Figure 6 caption, renumbered and re-pointed
    return caps


def blank_template() -> Document:
    tmp = OUT / "_t.docx"
    OUT.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TEMPLATE, tmp)
    doc = Document(str(tmp))
    for el in list(doc.element.body.iterchildren()):
        if el.tag != W + "sectPr":
            doc.element.body.remove(el)
    tmp.unlink()
    return doc


def build_figures():
    doc = blank_template()
    caps = captions()
    for n in range(1, 8):
        img = next(FIGDIR.glob("Figure%d_*.png" % n))
        title, body = DISP.split_title(DISP.tag_plain_subscripts(caps[n]))
        label, rest = title.split(".", 1)
        p = doc.add_paragraph()
        p.style = doc.styles["Normal"]
        p.paragraph_format.line_spacing = 1.15
        DISP.add_rich(p, "**%s.**%s" % (label, rest))
        im = Image.open(img)
        w, h = MAX_W_IN, MAX_W_IN * im.height / im.width
        if h > MAX_H_IN:
            h, w = MAX_H_IN, MAX_H_IN * im.width / im.height
        pic = doc.add_paragraph()
        pic.style = doc.styles["Normal"]
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(str(img), width=Inches(w))
        if body:
            c = doc.add_paragraph()
            c.style = doc.styles["Caption"]
            c.paragraph_format.line_spacing = None
            DISP.add_rich(c, body)
        if n < 7:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        print("  Figure %d  %-36s %.2f x %.2f in  caption %d words" % (n, img.name, w, h, len(body.split())))
    out = OUT / "FIGURE_maintext_draft02.docx"
    doc.save(str(out))
    return out


def build_tables():
    doc = blank_template()
    src = Document(str(BUILT_MAIN))
    kids = list(src.element.body.iterchildren())
    n = 0
    for i, el in enumerate(kids):
        if el.tag != W + "tbl":
            continue
        n += 1
        title = next(k for k in reversed(kids[:i]) if k.tag == W + "p" and Paragraph(k, src).text.strip())
        note = next(k for k in kids[i + 1:] if k.tag == W + "p" and Paragraph(k, src).text.strip())
        body = doc.element.body
        tp = copy.deepcopy(title)
        body.insert(len(body) - 1, tp)
        Paragraph(tp, doc).paragraph_format.line_spacing = 1.15
        body.insert(len(body) - 1, copy.deepcopy(el))
        np_ = copy.deepcopy(note)
        body.insert(len(body) - 1, np_)
        npar = Paragraph(np_, doc)
        npar.style = doc.styles["Caption"]
        npar.paragraph_format.line_spacing = None
        t = src.tables[n - 1]
        print("  Table %d  %dx%d  style=%s" % (n, len(t.rows), len(t.columns), t.style.name))
    stats = FIX.fix_tables(doc, [q for q in doc.paragraphs if q.style.name == "Caption"])
    left = FIX.bare_cl(doc)
    if left:
        raise SystemExit("bare 'CL' left: %s" % left)
    print("  terminology + subscripts: %d cells, %d notes" % (stats["cells"], stats["notes"]))
    out = OUT / "TABLE_maintext_draft02.docx"
    doc.save(str(out))
    return out


def main():
    print("figures:")
    f = build_figures()
    print("tables:")
    t = build_tables()
    for p in (f, t):
        d = Document(str(p))
        txt = "\n".join(q.text for q in d.paragraphs)
        bad = [m for m in ("<sub", "<sup", "**", "<<<") if m in txt]
        print("%-32s %d paragraphs, %d tables, %d images, %.0f kB%s"
              % (p.name, len(d.paragraphs), len(d.tables), len(d.inline_shapes), p.stat().st_size / 1e3,
                 "  LITERAL MARKUP %s" % bad if bad else ""))
        if bad:
            raise SystemExit("literal markup in %s" % p.name)


if __name__ == "__main__":
    main()
