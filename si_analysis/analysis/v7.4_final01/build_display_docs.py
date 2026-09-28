#!/usr/bin/env python3
"""v7.4 final_01 — paste-ready display documents: one for Figures, one for Tables.

Both use `sources/tables/TABLE_maintext.docx` as the template, so styles, fonts, page setup and section
properties are identical to the file the author already pastes from: 12 pt Times New Roman body, the item
title as a Normal paragraph at 1.15 line spacing with the label in bold, and the note/caption in the
`Caption` style (9 pt italic, colour 0E2841).

  FIGURE_maintext_v74.docx   title + figure + caption, Figures 1-6, one per page
  TABLE_maintext_v74.docx    title + table + note, Tables 1-3

Nothing is retyped: table objects and their title/note paragraphs are deep-copied from the built
`draft/v7.4_methods_results.docx` (so the corrected Table 2 cell and the revised Table 3 title and note come
along), and figure captions are split out of the frozen Round 4C package.

    .venv_pbpk/bin/python repo/claude_new_Sep21/analysis/v7.4_final01/build_display_docs.py
"""
from __future__ import annotations

import copy
import importlib.util
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from PIL import Image

S21 = Path(__file__).resolve().parents[2]
FINAL = S21 / "manuscript" / "v7.4" / "final_01"
TEMPLATE = FINAL / "sources" / "tables" / "TABLE_maintext.docx"
BUILT = FINAL / "draft" / "v7.4_methods_results.docx"
FIGDIR = FINAL / "sources" / "figures"
OUT = FINAL / "display"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

MAX_W_IN = 6.5   # text width: 8.5in page less two 1in margins
MAX_H_IN = 7.2   # leaves room for the title and caption on the same page


def load(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


MD = load(S21 / "manuscript" / "v4" / "scripts" / "md2docx.py", "md2docx")
BUILD = load(Path(__file__).with_name("build_final01.py"), "build_final01")
FIX = load(Path(__file__).with_name("table_fixes.py"), "table_fixes")

# Figure 1's caption is v7.2 paragraph 18 (Introduction, out of scope, title only, no caption body).
FIG1_TITLE = ("**FIGURE 1. Schematic overview of the study workflow, illustrating hybrid DL–ML prediction of "
              "pharmacokinetic parameters and their integration into downstream PBPK modeling and simulation.**")


def blank_template() -> Document:
    """A copy of TABLE_maintext.docx with an empty body but its styles and section properties intact."""
    tmp = OUT / "_template.docx"
    OUT.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TEMPLATE, tmp)
    doc = Document(str(tmp))
    body = doc.element.body
    for el in list(body.iterchildren()):
        if el.tag != W + "sectPr":
            body.remove(el)
    tmp.unlink()
    return doc


def add_rich(par, text, base_bold=False):
    """md2docx.add_runs handles markdown and unicode sub/superscripts, but the frozen captions use HTML
    <sub> tags (CL<sub>sys</sub> has no unicode subscript for 'y'), so those become real Word runs here."""
    for seg in re.split(r"(<sub>.*?</sub>|<sup>.*?</sup>)", text, flags=re.S):
        if not seg:
            continue
        m = re.fullmatch(r"<(sub|sup)>(.*?)</\1>", seg, re.S)
        if m:
            r = par.add_run(m.group(2))
            if m.group(1) == "sub":
                r.font.subscript = True
            else:
                r.font.superscript = True
            r.bold = True if base_bold else None
        else:
            MD.add_runs(par, seg, base_bold=base_bold)


def add_para(doc, text, style, line_spacing=None):
    p = doc.add_paragraph()
    p.style = doc.styles[style]
    p.paragraph_format.line_spacing = line_spacing
    add_rich(p, text)
    return p


def figure_captions() -> dict[int, str]:
    caps = {1: FIG1_TITLE}
    body = (FINAL / "src" / "body_results.md").read_text()
    m = re.search(r"^(\*\*FIGURE 2\..*?)$", body, re.M)
    caps[2] = m.group(1).strip()
    caps.update(BUILD.frozen_captions())          # Figures 3-6, verbatim from the Round 4C freeze
    return caps


# Figures 3-6 come from the frozen package, which already writes CL<sub>sys</sub> etc.; Figures 1 and 2 are
# hand-written and spell those tokens plainly. Tagging the plain ones here keeps all six consistent, and it
# goes through add_rich rather than a run rewrite, so the bold panel labels inside captions survive.
PLAIN_SUB = re.compile(r"\b(CLsys|CLint|VDss|Cmax|Fu|fu|Kp|pKa)\b")
SPLITS = {"CLsys": ("CL", "sys"), "CLint": ("CL", "int"), "VDss": ("VD", "ss"), "Cmax": ("C", "max"),
          "Fu": ("F", "u"), "fu": ("f", "u"), "Kp": ("K", "p"), "pKa": ("pK", "a")}


def tag_plain_subscripts(text):
    def rep(m):
        base, sub = SPLITS[m.group(1)]
        return "%s<sub>%s</sub>" % (base, sub)
    return PLAIN_SUB.sub(rep, text)


def split_title(text):
    """Leading bold span is the title; the remainder is the caption body."""
    m = re.match(r"\s*\*\*(.+?)\*\*\s*(.*)$", text, re.S)
    if not m:
        return text.strip(), ""
    return m.group(1).strip(), re.sub(r"\s+", " ", m.group(2)).strip()


def build_figures():
    doc = blank_template()
    caps = figure_captions()
    for n in range(1, 7):
        img = next(FIGDIR.glob("Figure%d_*.png" % n))
        title, body = split_title(tag_plain_subscripts(caps[n]))
        label, rest = title.split(".", 1)
        add_para(doc, "**%s.**%s" % (label, rest), "Normal", line_spacing=1.15)

        im = Image.open(img)
        w, h = MAX_W_IN, MAX_W_IN * im.height / im.width
        if h > MAX_H_IN:
            h, w = MAX_H_IN, MAX_H_IN * im.width / im.height
        pic = doc.add_paragraph()
        pic.style = doc.styles["Normal"]
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(str(img), width=Inches(w))

        if body:
            add_para(doc, body, "Caption")
        if n < 6:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        print("  Figure %d  %-34s %.2f x %.2f in   caption %d words"
              % (n, img.name, w, h, len(body.split())))
    out = OUT / "FIGURE_maintext_v74.docx"
    doc.save(str(out))
    return out


def build_tables():
    doc = blank_template()
    src = Document(str(BUILT))
    kids = list(src.element.body.iterchildren())
    n = 0
    for i, el in enumerate(kids):
        if el.tag != W + "tbl":
            continue
        n += 1
        title = note = None
        for j in range(i - 1, -1, -1):
            if kids[j].tag == W + "p" and Paragraph(kids[j], src).text.strip():
                title = kids[j]
                break
        for j in range(i + 1, len(kids)):
            if kids[j].tag == W + "p" and Paragraph(kids[j], src).text.strip():
                note = kids[j]
                break
        # title paragraph: copied verbatim so its bold spans are exactly as approved
        tp = copy.deepcopy(title)
        doc.element.body.insert(len(doc.element.body) - 1, tp)
        Paragraph(tp, doc).paragraph_format.line_spacing = 1.15
        doc.element.body.insert(len(doc.element.body) - 1, copy.deepcopy(el))
        np_ = copy.deepcopy(note)
        doc.element.body.insert(len(doc.element.body) - 1, np_)
        npar = Paragraph(np_, doc)
        npar.style = doc.styles["Caption"]   # match TABLE_maintext.docx
        npar.paragraph_format.line_spacing = None
        for _ in range(2):
            doc.element.body.insert(len(doc.element.body) - 1, copy.deepcopy(doc.add_paragraph()._p))
            doc.element.body.remove(doc.element.body[-2])
        t = src.tables[n - 1]
        print("  Table %d  %dx%d  style=%s  note %d words"
              % (n, len(t.rows), len(t.columns), t.style.name if t.style else None,
                 len(Paragraph(note, src).text.split())))
    notes = [q for q in doc.paragraphs if q.style.name == "Caption"]
    stats = FIX.fix_tables(doc, notes)
    print("  terminology + subscripts: %d cells, %d notes rewritten" % (stats["cells"], stats["notes"]))
    left = FIX.bare_cl(doc)
    if left:
        raise SystemExit("bare 'CL' left in the tables: %s" % left)
    out = OUT / "TABLE_maintext_v74.docx"
    doc.save(str(out))
    return out


def main():
    print("figures:")
    f = build_figures()
    print("tables:")
    t = build_tables()
    bad = []
    for p in (f, t):
        d = Document(str(p))
        txt = "\n".join(q.text for q in d.paragraphs)
        for pat in ("<sub", "</sub", "<sup", "</sup", "**", "<<<"):
            if pat in txt:
                bad.append((p.name, pat, txt.count(pat)))
        print("%-28s %d paragraphs, %d tables, %d images, %.0f kB"
              % (p.name, len(d.paragraphs), len(d.tables), len(d.inline_shapes), p.stat().st_size / 1e3))
    if bad:
        raise SystemExit("literal markup left in the output: %s" % bad)
    print("no literal markup left in either document")


if __name__ == "__main__":
    main()
