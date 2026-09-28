#!/usr/bin/env python3
"""v8.0 — restore the two-objective framing around FIGURE 1 and give the figure a real caption.

FIGURE 1 lays out OBJECTIVE 1 (1A development / 1B evaluation) and OBJECTIVE 2 (2A development /
2B evaluation). Two things were out of step with it:

  * the caption was one generic sentence that described none of that structure;
  * the paragraph that introduces the figure carried a comma splice ("platform., we investigated") and,
    in the Round 5C framing draft, I had deleted its (i)/(ii) enumeration as a duplicate of the two
    questions. With this figure present that deletion was wrong: paragraph 19 asks the two research
    QUESTIONS, and this paragraph states the two OBJECTIVES the figure is drawn around. They are not
    the same sentence and the figure needs its text counterpart.

No figure is added or removed, so no renumbering is involved: FIGURE 1 stays FIGURE 1.

    .venv_pbpk/bin/python repo/claude_new_Sep21/analysis/v7.4_final01/build_figure1_intro_v8.py
"""
from __future__ import annotations

import importlib.util
import re
import shutil
import sys
from pathlib import Path

from docx import Document

S21 = Path(__file__).resolve().parents[2]
D2 = S21 / "manuscript" / "v7.4" / "final_01" / "draft_02"
SRC = D2 / "v8_framing" / "v8.0_abstract_conclusions.docx"
OUT = D2 / "v8_framing"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
BD = None

# The caption walks the four panels in the order the figure draws them. It describes the figure and
# introduces no result. "two clearance parameterizations" is deliberately neutral: the figure art still
# carries the labels "Bottom-up / Top-down", so this wording is correct whether or not that art is updated.
CAPTION = (
    "FIGURE 1. Overview of the two-objective study design. Objective 1, representation benchmarking. "
    "(1A) Development of the DL–ML pipeline: SMILES inputs are encoded as molecular graphs by a graph "
    "neural network, and the learned embeddings, alone or merged with RDKit descriptors, are passed to "
    "conventional regressors (support vector machine, random forest, XGBoost) to predict Fu, "
    "CL<sub>sys</sub> and VD<sub>ss</sub>. (1B) Evaluation on the held-out Test #1 benchmark using R<sup>2</sup>, "
    "RMSE, MAE, GMFE and the fraction of predictions within twofold, which gives the representation "
    "comparison. Objective 2, hybrid PBPK modeling, carries the structure-derived PK parameters forward as "
    "simulation inputs. (2A) Development of the hybrid workflow in Simcyp with R-based post-processing: "
    "predicted parameters are integrated into compound files and compared across model structures, minimal "
    "PBPK under two clearance parameterizations against full PBPK under three tissue-distribution methods. "
    "(2B) Evaluation against clinical concentration–time profiles for the exclusive Test #2 compounds, "
    "scored by signed and absolute log<sub>2</sub> fold error in AUC<sub>0–t</sub> and C<sub>max</sub> and "
    "by concentration–time profile error."
)

# splice repaired, and the two objectives restated as objectives rather than as a second copy of the
# questions in paragraph 19
INTRO = (
    "In this study, as shown in Figure 1, we developed a graph attention-based deep learning-encoder plus "
    "machine learning-regressor (DL–ML) framework to predict key human PK parameters and evaluated its "
    "integration into mechanistic PBPK simulations within the Simcyp platform. The work is organized around "
    "two objectives. The first is representation benchmarking, in which models built on merged graph-derived "
    "and RDKit descriptors are compared with descriptor-based models on a held-out benchmark set (Test #1). "
    "The second is hybrid PBPK modeling, in which the resulting structure-derived parameters are supplied to "
    "PBPK simulations and evaluated against clinical concentration–time profiles for an exclusive set of "
    "compounds (Test #2)."
)


def build(highlight: bool):
    out = OUT / ("v8.0_framing_full_highlighted.docx" if highlight else "v8.0_framing_full.docx")
    shutil.copy2(SRC, out)
    doc = Document(str(out))
    ps = doc.paragraphs
    mode = "diff" if highlight else "clean"

    cap = next(i for i, p in enumerate(ps) if p.text.strip().startswith("FIGURE 1."))
    intro = next(i for i, p in enumerate(ps) if "as shown in Figure 1" in p.text)
    if not (cap < intro):
        raise SystemExit("expected the caption before the introducing paragraph; got %d / %d" % (cap, intro))

    # the figure image is INLINE in the caption paragraph, so the caption text is rewritten run-by-run
    # rather than by clearing the paragraph, which would delete the image with it
    img = len(ps[cap]._p.findall(".//" + W + "drawing"))
    if img != 1:
        raise SystemExit("expected exactly one inline image in the caption paragraph, found %d" % img)
    if not BD.rich_replace(ps[cap], ps[cap].text.strip(), CAPTION, mode):
        raise SystemExit("caption anchor not found")
    if len(ps[cap]._p.findall(".//" + W + "drawing")) != 1:
        raise SystemExit("the inline image was lost while rewriting the caption")

    old = ps[intro].text
    if not BD.rich_replace(ps[intro], old, INTRO, mode):
        raise SystemExit("intro anchor not found")

    doc.save(str(out))
    return out, cap, intro


def main():
    global BD
    spec = importlib.util.spec_from_file_location(
        "bd", Path(__file__).with_name("build_discussion_round5c.py"))
    BD = importlib.util.module_from_spec(spec)
    sys.modules["bd"] = BD
    spec.loader.exec_module(BD)
    s2 = importlib.util.spec_from_file_location(
        "hl", S21 / "manuscript" / "v4" / "scripts" / "build_highlighted_v4.py")
    BD.HL = importlib.util.module_from_spec(s2)
    s2.loader.exec_module(BD.HL)

    for hi in (False, True):
        out, cap, intro = build(hi)
        print("%-44s caption ¶%d, intro ¶%d" % (out.name, cap, intro))

    d = Document(str(OUT / "v8.0_framing_full.docx"))
    c = next(p for p in d.paragraphs if p.text.strip().startswith("FIGURE 1."))
    i = next(p for p in d.paragraphs if "as shown in Figure 1" in p.text)
    print("\ncaption: %d words, image preserved: %s"
          % (len(c.text.split()), len(c._p.findall(".//" + W + "drawing")) == 1))
    print("intro  : %d words" % len(i.text.split()))


if __name__ == "__main__":
    main()
